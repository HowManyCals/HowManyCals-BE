from __future__ import annotations

import csv
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, Iterator

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SEED_DIR = ROOT / "src" / "main" / "resources" / "seed"
OUTPUT_DIR = SEED_DIR / "preprocessed"
AI_INPUT = SEED_DIR / "음식분류_AI_데이터_영양DB.xlsx"
PACKAGED_INPUT = SEED_DIR / "20260429_가공식품_277074건.xlsx"

AI_OUTPUT = OUTPUT_DIR / "ai_food_entry_seed.csv"
AI_ALIAS_OUTPUT = OUTPUT_DIR / "ai_food_entry_alias_seed.csv"
PACKAGED_OUTPUT = OUTPUT_DIR / "packaged_food_seed.csv"
PACKAGED_ALIAS_OUTPUT = OUTPUT_DIR / "packaged_food_alias_seed.csv"
REPORT_OUTPUT = OUTPUT_DIR / "preprocess_report.md"

CORPORATE_STOPWORDS = {
    "주식회사", "(주)", "주", "유한회사", "합자회사", "재단법인", "사단법인",
    "co", "ltd", "inc", "corp", "company", "limited"
}
MARKETING_STOPWORDS = {
    "big", "mini", "오리지널", "original", "프리미엄", "premium", "리얼", "real",
    "점보", "jumbo", "대용량", "한입", "extra", "더블", "트리플", "classic",
    "best", "choice", "chef", "yeschef"
}
ATTRIBUTE_STOPWORDS = {
    "우리밀", "국내산", "유기농", "무농약", "자연산", "국산"
}
GENERIC_REPRESENTATIVE_NAMES = {
    "과자", "떡", "스낵과자", "음료", "빵", "라면", "면", "떡류", "가공식품", "해당없음"
}

HEADER_NORMALIZATION_PATTERN = re.compile(r"[\s\u3000]+")
NON_WORD_KEY_PATTERN = re.compile(r"[^0-9a-z가-힣]+")
TOKEN_SPLIT_PATTERN = re.compile(r"[\s/+,·|()\[\]{}<>'\"`~!@#$%^&*=?:;._-]+")
ASCII_TOKEN_PATTERN = re.compile(r"^[a-z0-9]+$")
KOREAN_PATTERN = re.compile(r"[가-힣]")
WEIGHT_TOKEN_PATTERN = re.compile(r"^\d+(?:\.\d+)?(?:g|kg|ml|l|개|입|봉|캔|병)$", re.IGNORECASE)


@dataclass
class AiFoodEntryRow:
    ai_food_name: str
    normalized_ai_food_name: str
    serving_weight: str
    serving_unit: str
    serving_kcal: str
    carbohydrate: str
    protein: str
    fat: str
    source_sheet: str


@dataclass
class PackagedFoodRow:
    product_code: str
    product_name_raw: str
    normalized_product_name_raw: str
    manufacturer_name: str
    representative_food_name: str
    normalized_representative_food_name: str
    canonical_product_name: str
    normalized_canonical_product_name: str
    canonical_confidence: str
    serving_reference: str
    package_weight: str
    kcal: str
    carbohydrate: str
    protein: str
    fat: str
    sugar: str
    category_large: str
    category_medium: str
    category_small: str
    source_name: str


@dataclass(frozen=True)
class AliasRow:
    source_code: str
    alias_name: str
    normalized_alias: str
    alias_type: str
    priority: int


def normalize_text(value: object) -> str:
    if value is None:
        return ""
    text = str(value).replace("\ufeff", "").replace("\u3000", " ")
    text = HEADER_NORMALIZATION_PATTERN.sub(" ", text).strip()
    return text


def normalize_header(value: object) -> str:
    return normalize_text(value).replace(" ", "")


def normalize_key(value: object) -> str:
    text = normalize_text(value).lower()
    if not text:
        return ""
    text = NON_WORD_KEY_PATTERN.sub("", text)
    return text


def split_tokens(value: str) -> list[str]:
    return [token for token in TOKEN_SPLIT_PATTERN.split(normalize_text(value)) if token]


def parse_float_string(value: object) -> str:
    text = normalize_text(value)
    if not text or text == "-":
        return ""
    try:
        number = float(text)
    except ValueError:
        return text
    return str(int(number)) if number.is_integer() else f"{number:.4f}".rstrip("0").rstrip(".")


def manufacturer_tokens(value: str) -> set[str]:
    tokens = {normalize_key(token) for token in split_tokens(value)}
    return {token for token in tokens if token and token not in CORPORATE_STOPWORDS}


def contains_korean(value: str) -> bool:
    return bool(KOREAN_PATTERN.search(value))


def build_packaged_canonical_name(product_name: str, representative_name: str, manufacturer_name: str) -> tuple[str, str, list[str]]:
    raw_tokens = split_tokens(product_name)
    if not raw_tokens:
        return "", "LOW", []

    manufacturer_token_set = manufacturer_tokens(manufacturer_name)
    representative_tokens = [token for token in split_tokens(representative_name) if normalize_key(token)]
    representative_token_keys = {normalize_key(token) for token in representative_tokens}
    representative_key = normalize_key(representative_name)
    has_korean_token = any(contains_korean(token) for token in raw_tokens)

    filtered_tokens: list[str] = []
    removed_tokens: list[str] = []

    for token in raw_tokens:
        normalized_token = normalize_key(token)
        if not normalized_token:
            continue
        if normalized_token in manufacturer_token_set:
            removed_tokens.append(token)
            continue
        if normalized_token in MARKETING_STOPWORDS:
            removed_tokens.append(token)
            continue
        if WEIGHT_TOKEN_PATTERN.match(token.lower()):
            removed_tokens.append(token)
            continue
        if has_korean_token and ASCII_TOKEN_PATTERN.fullmatch(normalized_token) and normalized_token not in representative_token_keys:
            removed_tokens.append(token)
            continue
        filtered_tokens.append(token)

    if not filtered_tokens:
        filtered_tokens = raw_tokens[:]

    candidate_tokens = [
        token for token in filtered_tokens
        if normalize_key(token) not in ATTRIBUTE_STOPWORDS
    ]
    if not candidate_tokens:
        candidate_tokens = filtered_tokens[:]

    canonical_name = "".join(candidate_tokens)
    confidence = "MEDIUM"

    if representative_key and representative_key not in GENERIC_REPRESENTATIVE_NAMES:
        representative_display = "".join(representative_tokens) if representative_tokens else normalize_text(representative_name)
        modifier_tokens = [
            token for token in candidate_tokens
            if normalize_key(token) not in representative_token_keys
        ]
        if modifier_tokens:
            canonical_name = "".join(modifier_tokens) + representative_display
            confidence = "HIGH"
        else:
            canonical_name = representative_display
            confidence = "MEDIUM"
    elif removed_tokens:
        confidence = "MEDIUM"
    else:
        confidence = "LOW"

    canonical_name = normalize_text(canonical_name)
    if not canonical_name:
        canonical_name = normalize_text(product_name)
        confidence = "LOW"

    alias_candidates = build_packaged_alias_candidates(
        product_name=product_name,
        representative_name=representative_name,
        canonical_name=canonical_name,
        candidate_tokens=candidate_tokens,
        filtered_tokens=filtered_tokens,
    )
    return canonical_name, confidence, alias_candidates


def build_packaged_alias_candidates(
    product_name: str,
    representative_name: str,
    canonical_name: str,
    candidate_tokens: list[str],
    filtered_tokens: list[str],
) -> list[str]:
    aliases: list[str] = []

    def add(alias: str) -> None:
        alias = normalize_text(alias)
        if alias and alias not in aliases:
            aliases.append(alias)

    add(product_name)
    add(canonical_name)

    representative_clean = normalize_text(representative_name)
    representative_key = normalize_key(representative_clean)
    if representative_clean and representative_key not in GENERIC_REPRESENTATIVE_NAMES:
        add(representative_clean)

    if candidate_tokens:
        add("".join(candidate_tokens))
        add(" ".join(candidate_tokens))

    if filtered_tokens:
        add("".join(filtered_tokens))

    if representative_clean and representative_key not in GENERIC_REPRESENTATIVE_NAMES:
        rep_token_keys = {normalize_key(token) for token in split_tokens(representative_clean)}
        modifiers = [token for token in candidate_tokens if normalize_key(token) not in rep_token_keys]
        if modifiers:
            add("".join(modifiers) + representative_clean)
            add(" ".join(modifiers + [representative_clean]))

    return aliases


def iter_worksheet_rows(path: Path) -> tuple[str, Iterator[tuple[object, ...]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet_name = workbook.sheetnames[0]
    worksheet = workbook[sheet_name]
    rows = worksheet.iter_rows(values_only=True)
    return sheet_name, rows


def preprocess_ai_food_entries() -> dict[str, int]:
    sheet_name, rows = iter_worksheet_rows(AI_INPUT)
    header = next(rows)
    index_by_header = {normalize_header(column): idx for idx, column in enumerate(header)}

    required_headers = {
        "음식명": "food_name",
        "중량(g)": "weight",
        "에너지(kcal)": "kcal",
        "탄수화물(g)": "carbohydrate",
        "지방(g)": "fat",
        "단백질(g)": "protein",
    }
    missing_headers = [header_name for header_name in required_headers if header_name not in index_by_header]
    if missing_headers:
        raise KeyError(f"AI 영양 DB 필수 헤더 누락: {missing_headers}")

    best_rows: dict[str, AiFoodEntryRow] = {}
    duplicate_count = 0

    for row in rows:
        food_name = normalize_text(row[index_by_header["음식명"]])
        normalized_food_name = normalize_key(food_name)
        if not normalized_food_name:
            continue

        candidate_row = AiFoodEntryRow(
            ai_food_name=food_name,
            normalized_ai_food_name=normalized_food_name,
            serving_weight=parse_float_string(row[index_by_header["중량(g)"]]),
            serving_unit="g",
            serving_kcal=parse_float_string(row[index_by_header["에너지(kcal)"]]),
            carbohydrate=parse_float_string(row[index_by_header["탄수화물(g)"]]),
            protein=parse_float_string(row[index_by_header["단백질(g)"]]),
            fat=parse_float_string(row[index_by_header["지방(g)"]]),
            source_sheet=sheet_name,
        )

        existing = best_rows.get(normalized_food_name)
        if existing is None:
            best_rows[normalized_food_name] = candidate_row
            continue

        duplicate_count += 1
        if score_ai_row(candidate_row) > score_ai_row(existing):
            best_rows[normalized_food_name] = candidate_row

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    with AI_OUTPUT.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=list(AiFoodEntryRow.__annotations__.keys()))
        writer.writeheader()
        for row in sorted(best_rows.values(), key=lambda item: item.normalized_ai_food_name):
            writer.writerow(row.__dict__)

    with AI_ALIAS_OUTPUT.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=["source_code", "alias_name", "normalized_alias", "alias_type", "priority"])
        writer.writeheader()
        for row in sorted(best_rows.values(), key=lambda item: item.normalized_ai_food_name):
            aliases = [
                AliasRow(row.normalized_ai_food_name, row.ai_food_name, row.normalized_ai_food_name, "AI_NAME", 100),
            ]
            for alias in aliases:
                writer.writerow(alias.__dict__)

    return {
        "row_count": len(best_rows),
        "duplicate_count": duplicate_count,
    }


def score_ai_row(row: AiFoodEntryRow) -> int:
    score = 0
    for value in (row.serving_weight, row.serving_kcal, row.carbohydrate, row.protein, row.fat):
        if value:
            score += 1
    return score


def preprocess_packaged_foods() -> dict[str, int]:
    sheet_name, rows = iter_worksheet_rows(PACKAGED_INPUT)
    header = next(rows)
    index_by_header = {normalize_text(column): idx for idx, column in enumerate(header)}

    required_headers = [
        "식품코드", "식품명", "대표식품명", "제조사명", "영양성분함량기준량", "식품중량",
        "에너지(kcal)", "탄수화물(g)", "단백질(g)", "지방(g)", "당류(g)",
        "식품대분류명", "식품중분류명", "식품소분류명", "출처명"
    ]
    missing_headers = [header_name for header_name in required_headers if header_name not in index_by_header]
    if missing_headers:
        raise KeyError(f"가공식품 DB 필수 헤더 누락: {missing_headers}")

    row_count = 0
    alias_count = 0
    canonical_confidence_counts = {"HIGH": 0, "MEDIUM": 0, "LOW": 0}

    with PACKAGED_OUTPUT.open("w", encoding="utf-8-sig", newline="") as food_file, \
            PACKAGED_ALIAS_OUTPUT.open("w", encoding="utf-8-sig", newline="") as alias_file:
        food_writer = csv.DictWriter(food_file, fieldnames=list(PackagedFoodRow.__annotations__.keys()))
        alias_writer = csv.DictWriter(alias_file, fieldnames=["source_code", "alias_name", "normalized_alias", "alias_type", "priority"])
        food_writer.writeheader()
        alias_writer.writeheader()

        for row in rows:
            product_code = normalize_text(row[index_by_header["식품코드"]])
            product_name = normalize_text(row[index_by_header["식품명"]])
            if not product_code or not product_name:
                continue

            if row_count > 0 and row_count % 25000 == 0:
                print(f"[packaged] processed {row_count} rows...")

            representative_name = normalize_text(row[index_by_header["대표식품명"]])
            manufacturer_name = normalize_text(row[index_by_header["제조사명"]])
            canonical_name, confidence, alias_candidates = build_packaged_canonical_name(
                product_name=product_name,
                representative_name=representative_name,
                manufacturer_name=manufacturer_name,
            )

            packaged_row = PackagedFoodRow(
                product_code=product_code,
                product_name_raw=product_name,
                normalized_product_name_raw=normalize_key(product_name),
                manufacturer_name=manufacturer_name,
                representative_food_name=representative_name,
                normalized_representative_food_name=normalize_key(representative_name),
                canonical_product_name=canonical_name,
                normalized_canonical_product_name=normalize_key(canonical_name),
                canonical_confidence=confidence,
                serving_reference=normalize_text(row[index_by_header["영양성분함량기준량"]]),
                package_weight=normalize_text(row[index_by_header["식품중량"]]),
                kcal=parse_float_string(row[index_by_header["에너지(kcal)"]]),
                carbohydrate=parse_float_string(row[index_by_header["탄수화물(g)"]]),
                protein=parse_float_string(row[index_by_header["단백질(g)"]]),
                fat=parse_float_string(row[index_by_header["지방(g)"]]),
                sugar=parse_float_string(row[index_by_header["당류(g)"]]),
                category_large=normalize_text(row[index_by_header["식품대분류명"]]),
                category_medium=normalize_text(row[index_by_header["식품중분류명"]]),
                category_small=normalize_text(row[index_by_header["식품소분류명"]]),
                source_name=normalize_text(row[index_by_header["출처명"]]),
            )
            food_writer.writerow(packaged_row.__dict__)
            row_count += 1
            canonical_confidence_counts[confidence] = canonical_confidence_counts.get(confidence, 0) + 1

            for alias in build_packaged_alias_rows(product_code, product_name, canonical_name, representative_name, alias_candidates):
                alias_writer.writerow(alias.__dict__)
                alias_count += 1

    return {
        "row_count": row_count,
        "alias_count": alias_count,
        "high_confidence_count": canonical_confidence_counts["HIGH"],
        "medium_confidence_count": canonical_confidence_counts["MEDIUM"],
        "low_confidence_count": canonical_confidence_counts["LOW"],
        "sheet_name": sheet_name,
    }


def build_packaged_alias_rows(
    product_code: str,
    product_name: str,
    canonical_name: str,
    representative_name: str,
    alias_candidates: list[str],
) -> list[AliasRow]:
    alias_rows: list[AliasRow] = []
    seen: set[str] = set()

    def add(alias_name: str, alias_type: str, priority: int) -> None:
        normalized_alias = normalize_key(alias_name)
        if not normalized_alias or normalized_alias in seen:
            return
        seen.add(normalized_alias)
        alias_rows.append(AliasRow(product_code, normalize_text(alias_name), normalized_alias, alias_type, priority))

    add(product_name, "RAW_PRODUCT_NAME", 100)
    add(canonical_name, "CANONICAL_PRODUCT_NAME", 95)

    for alias_name in alias_candidates:
        alias_type = "CANONICAL_VARIANT" if normalize_key(alias_name) != normalize_key(product_name) else "RAW_VARIANT"
        priority = 90 if alias_type == "CANONICAL_VARIANT" else 80
        add(alias_name, alias_type, priority)

    representative_key = normalize_key(representative_name)
    if representative_name and representative_key and representative_key not in GENERIC_REPRESENTATIVE_NAMES:
        add(representative_name, "REPRESENTATIVE_NAME", 60)

    return alias_rows


def write_report(ai_stats: dict[str, int], packaged_stats: dict[str, int]) -> None:
    report = f"""# Seed preprocess report

## AI nutrition DB
- deduplicated rows: {ai_stats['row_count']}
- duplicates removed: {ai_stats['duplicate_count']}
- output: `{AI_OUTPUT.relative_to(ROOT)}`
- alias output: `{AI_ALIAS_OUTPUT.relative_to(ROOT)}`

## Packaged food DB
- rows processed: {packaged_stats['row_count']}
- alias rows generated: {packaged_stats['alias_count']}
- canonical confidence HIGH: {packaged_stats['high_confidence_count']}
- canonical confidence MEDIUM: {packaged_stats['medium_confidence_count']}
- canonical confidence LOW: {packaged_stats['low_confidence_count']}
- source sheet: {packaged_stats['sheet_name']}
- output: `{PACKAGED_OUTPUT.relative_to(ROOT)}`
- alias output: `{PACKAGED_ALIAS_OUTPUT.relative_to(ROOT)}`

## Notes
- `ai_food_entry_seed.csv`는 AI 고신뢰 exact source 후보용 seed입니다.
- `packaged_food_seed.csv`는 원본 상품명, 대표식품명, 제조사명과 함께 canonical 상품명을 생성합니다.
- `packaged_food_alias_seed.csv`는 exact index 후보를 만들기 위한 다중 alias seed입니다.
- canonical confidence가 LOW인 데이터는 서비스 exact 인덱스 노출 전 추가 검토가 필요합니다.
"""
    REPORT_OUTPUT.write_text(report, encoding="utf-8")


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    ai_stats = preprocess_ai_food_entries()
    packaged_stats = preprocess_packaged_foods()
    write_report(ai_stats, packaged_stats)
    print("AI rows:", ai_stats)
    print("Packaged rows:", packaged_stats)
    print("Outputs written under:", OUTPUT_DIR)


if __name__ == "__main__":
    main()


